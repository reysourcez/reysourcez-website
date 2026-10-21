#!/usr/bin/env python3
# Market Radar - settings workbook generator v1.7.0 (2026-10-13)
"""
Writes market-radar-settings-reference.xlsx by READING the shipped files, so the sheet cannot drift from the code
(a lesson this project learned the hard way). Run from the folder that holds the Market Radar files:

    python tools/market-radar-build-settings.py

Needs: Python 3, openpyxl, and Node (to read the engine's tables). If a value cannot be found in the code the
script stops with an error instead of writing a guess.
"""
import importlib.util
import sys
sys.dont_write_bytecode = True   # importing the builder must not leave a __pycache__ folder behind
import json
import os
import re
import subprocess
import sys
import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
VERSION = '1.7.0'
STAMP = time.strftime('%Y-%m-%d')
OUT = os.path.join(ROOT, 'market-radar-settings-reference.xlsx')

FONT = 'Arial'
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True, color='1F6F5C')
F_HEAD = Font(name=FONT, size=10, bold=True, color='FFFFFF')
F_NOTE = Font(name=FONT, size=9, italic=True, color='5C6D64')
FILL_HEAD = PatternFill('solid', fgColor='1F6F5C')
FILL_YOU = PatternFill('solid', fgColor='FFF2CC')
FILL_GEN = PatternFill('solid', fgColor='EEF3F0')
SIDE = Side(style='thin', color='C9D3CE')
BORDER = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)
WRAP = Alignment(wrap_text=True, vertical='top')


def read(rel):
    with open(os.path.join(ROOT, rel), 'r', encoding='utf-8') as fh:
        return fh.read()


def grab(src, pattern, what, flags=0):
    m = re.search(pattern, src, flags)
    if not m:
        raise SystemExit('Could not find %s in the code (pattern %r). The code changed shape: update this generator, do not guess.' % (what, pattern))
    return m


def engine_dump():
    js = (
        "const E=require(process.argv[1]);"
        "const pick=(c)=>({key:c.key,label:c.label,group:c.group,needTier:c.needTier,need:c.need,crowded:c.crowded,capital:c.capital,"
        "osm:c.osm.map(p=>p.join('=')),msic:c.msic,iowrt:c.iowrt,geo:c.geo,ovt:c.ovt,founder:c.founder,niches:c.niches});"
        "console.log(JSON.stringify({version:E.VERSION,cats:E.CATEGORIES.map(pick),groups:E.GROUPS,"
        "anchors:E.ANCHORS.map(a=>({key:a.key,label:a.label,osm:a.osm.map(p=>p.join('=')),ovt:a.ovt})),"
        "anchorSize:E.ANCHOR_SIZE,anchorAffinity:E.ANCHOR_AFFINITY,pairs:E.AFFINITY_PAIRS,"
        "niches:E.NICHES.map(n=>({key:n.key,label:n.label,osm:n.osm})),arch:E.ARCH,archOrder:E.ARCH_PRIORITY,"
        "T:E.T,opp:E.OPP_WEIGHTS,comp:E.COMP_WEIGHTS,idea:E.IDEA_WEIGHTS,capital:E.CAPITAL_FACTOR,need:E.NEED_TIER,"
        "popNorm:E.POP_NORM,incomeNorm:E.INCOME_NORM,drinkHints:E.DRINK_HINTS}));"
    )
    res = subprocess.run(['node', '-e', js, os.path.join(ROOT, 'market-radar-engine.js')], capture_output=True, text=True)
    if res.returncode != 0:
        raise SystemExit('node could not read the engine: ' + res.stderr)
    return json.loads(res.stdout)


def load_builder():
    spec = importlib.util.spec_from_file_location('mrb', os.path.join(HERE, 'market-radar-build-overture.py'))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ---------------------------------------------------------------- sheet helpers
def title(ws, text, sub):
    ws['A1'] = text
    ws['A1'].font = F_TITLE
    ws['A2'] = sub
    ws['A2'].font = F_NOTE
    ws['A2'].alignment = Alignment(wrap_text=False)


def table(ws, headers, rows, widths, you_cols=(), start=4, wrap_cols=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start, column=j, value=h)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, row in enumerate(rows, start + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border, c.alignment = F_BASE, BORDER, WRAP
            c.fill = FILL_YOU if j in you_cols else FILL_GEN if j == 2 and 'Current' in headers[1] else PatternFill()
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    return start + len(rows)


def build():
    eng = engine_dump()
    if eng['version'] != VERSION:
        raise SystemExit('Engine says v%s but this generator is for v%s - update both together.' % (eng['version'], VERSION))
    page, worker, html = read('market-radar.js'), read('market-radar-proxy-worker.js'), read('market-radar.html')
    mrb = load_builder()

    # ---- values read out of the client
    towns = re.findall(r"(\w+): \{ label: '([^']+)', lat: ([\d.]+), lng: ([\d.]+), district: '([^']+)' \}", grab(page, r'const TOWNS = \{(.*?)\n  \};', 'TOWNS', re.S).group(1))
    catch = re.findall(r"(\w+): \{ label: '([^']+)', profile: '([^']+)', seconds: (\d+), fallbackM: (\d+) \}", grab(page, r'const CATCHMENTS = \{(.*?)\n  \};', 'CATCHMENTS', re.S).group(1))
    if len(towns) != 4 or len(catch) != 6:
        raise SystemExit('Expected 4 towns and 6 catchments in market-radar.js, found %d and %d' % (len(towns), len(catch)))
    m = grab(page, r"const MAX_PER_DAY = (\d+), USAGE_KEY = '([^']+)', MAX_TICKED = (\d+), MAX_CUSTOM = (\d+);", 'daily/ticked/custom limits')
    max_day, usage_key, max_ticked, max_custom = m.group(1), m.group(2), m.group(3), m.group(4)
    direct = re.findall(r"'(https://[^']+)'", grab(page, r'const DIRECT_MIRRORS = \[(.*?)\];', 'DIRECT_MIRRORS', re.S).group(1))
    ovt_path = grab(page, r"const OVT_PATH = '([^']+)';", 'OVT_PATH').group(1)
    worker_url = grab(page, r"const WORKER = '([^']+)';", 'WORKER').group(1)
    find_gap = grab(page, r'lastFind < (\d+)', 'place-search gap').group(1)
    ovp_s = grab(page, r"cm\.profile === 'driving-car' \? (\d+) : (\d+)", 'Overpass timeouts')
    page_ver = grab(page, r"const VERSION = '([\d.]+)'", 'page version').group(1)

    # ---- values read out of the Worker
    w_ver = grab(worker, r"const WORKER_VERSION = '([\d.]+)'", 'worker version').group(1)
    mirrors = re.findall(r"'(https://[^']+)'", grab(worker, r'const OVERPASS_MIRRORS = \[(.*?)\];', 'OVERPASS_MIRRORS', re.S).group(1))
    origins = re.findall(r"'(https://[^']+)'", grab(worker, r'const ALLOWED_ORIGINS = \[(.*?)\];', 'ALLOWED_ORIGINS', re.S).group(1))
    ext_ms = grab(worker, r'const EXTERNAL_CALL_TIMEOUT_MS = (\d+)', 'timeout').group(1)
    max_pairs = grab(worker, r'const MAX_TAG_PAIRS = (\d+)', 'MAX_TAG_PAIRS').group(1)
    max_span = grab(worker, r'const MAX_SPAN_DEG = ([\d.]+)', 'MAX_SPAN_DEG').group(1)
    gemini = grab(worker, r"const GEMINI_MODEL = '([^']+)'", 'GEMINI_MODEL').group(1)
    wai = grab(worker, r"const WORKERS_AI_MODEL = '([^']+)'", 'WORKERS_AI_MODEL').group(1)
    pop = re.findall(r'(\w+): (\d+),', grab(worker, r'const DISTRICT_POPULATION_FALLBACK_2023 = \{(.*?)\};', 'population fallback', re.S).group(1))
    ttl_ovp = grab(worker, r"'Cache-Control': 'public, max-age=(\d+)', 'X-MR-Source'", 'Overpass cache').group(1)
    ttl_iso = grab(worker, r"max-age=(\d+)' \} \}\);\s*await caches\.default\.put\(\s*cacheKey|public, max-age=(2592000)", 'isochrone cache').group(0)
    ttl_iso = re.search(r'(2592000)', ttl_iso).group(1)
    ttl_dem = grab(worker, r'medianIncome: out\.medianIncome \}, (\d+)\)', 'demographics cache').group(1)
    env_names = sorted(set(re.findall(r'env\.([A-Z][A-Z_]+)', worker)))
    scripts = re.findall(r'<script src="([^"]+)"', html)
    styles = re.findall(r'<link rel="stylesheet" href="([^"]+)"', html)

    wb = Workbook()

    # ================= ReadMe
    ws = wb.active
    ws.title = 'ReadMe'
    title(ws, 'Market Radar - Settings Reference - v%s' % VERSION, 'Generated %s from the shipped files by tools/market-radar-build-settings.py. If this sheet and the code ever disagree, the code is right.' % STAMP)
    lines = [
        ('What this is', 'A human-readable list of every number and name you might want to change, with where it lives. Grey cells are read out of the code. Yellow cells are yours: type what you would like instead.'),
        ('How to use it', '1) Find the setting. 2) Type your wished value in the yellow cell. 3) Give the sheet to the AI assistant (or edit the file at the "Where" location yourself). 4) Redeploy using MARKET_RADAR_SETUP_AND_GLOSSARY.md. Never put real API keys in this file.'),
        ('Judgement vs fact', 'Anything labelled assumption is the author\'s judgement, not a measurement: crowding densities, need tiers, capital levels, anchor and affinity weights, the thresholds in Settings. Change them if your local knowledge says so.'),
        ('Files and versions', 'All project files use the same version number so a mismatch is obvious. Stable file names are deliberate: other pages link to market-radar.html.'),
    ]
    r = 4
    for k, v in lines:
        ws.cell(row=r, column=1, value=k).font = F_BOLD
        c = ws.cell(row=r, column=2, value=v)
        c.font, c.alignment = F_BASE, WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 44
        r += 1
    r += 1
    heads = ['File', 'Version', 'Lives at', 'Deployed by', 'Notes']
    files = [
        ('market-radar.html', page_ver, 'GitHub Pages, site root', 'Upload to GitHub', 'Cache-buster on the script tags: change ?v=... on every release.'),
        ('market-radar.js', page_ver, 'GitHub Pages, site root', 'Upload to GitHub', 'State, network, rendering. No scoring logic.'),
        ('market-radar-engine.js', eng['version'], 'GitHub Pages, site root', 'Upload to GitHub', 'All scoring and classification (pure logic, fully unit-tested). Must be uploaded together with market-radar.js.'),
        ('market-radar-proxy-worker.js', w_ver, 'Cloudflare Worker "market-radar-proxy"', 'Paste into Edit code, Deploy', 'Shown in the Area snapshot tab after an analysis.'),
        ('market-radar-data/overture-<town>.json', '(built)', 'GitHub Pages, folder market-radar-data/', 'Upload to GitHub', 'Optional second data source. Built by tools/market-radar-build-overture.py on your PC. Without it the page uses OpenStreetMap only and says so.'),
        ('tools/market-radar-build-overture.py', VERSION, 'Your PC', 'Run it', 'python -m pip install overturemaps, then python market-radar-build-overture.py --town all'),
        ('tools/market-radar-build-settings.py', VERSION, 'Your PC / AI session', 'Run it', 'Regenerates THIS workbook from the code.'),
        ('tools/market-radar-build-handoff.py / market-radar-build-package.py', VERSION, 'AI session', 'Run it', 'Regenerate the handoff dossier / the _MANIFEST.md and the ALL-FILES zip.'),
        ('market-radar-v1.7.0-ALL-FILES.zip', VERSION, 'Download', 'Unzip', 'Every file below with folders kept; START_HERE.md is the first thing to read.'),
        ('START_HERE.md', VERSION, 'Project folder', 'n/a', 'One page: which file goes where.'),
        ('archive-v1.6.2/', 'v1.6.2 (old)', 'Project folder', 'n/a', 'Old page, script, docs and workbook; file names carry -v1.6.2. Only for rolling back (needs the old Worker too).'),
        ('MARKET_RADAR_HANDOFF.md', VERSION, 'Project folder', 'n/a', 'AI continuity dossier: give this to a new AI session first.'),
        ('MARKET_RADAR_SETUP_AND_GLOSSARY.md', VERSION, 'Project folder', 'n/a', 'Deploy steps, what was verified, jargon.'),
        ('MARKET_RADAR_HISTORY.md', 'archive (ends at v1.6.2)', 'Project folder', 'n/a', 'The old story up to v1.6.2, kept verbatim. Not out of date: it is an archive.'),
        ('tests/', VERSION, 'Project folder', 'n/a', 'engine.test.cjs, worker.test.mjs, overture-builder.test.py, e2e.mjs (+ run-all.sh).'),
    ]
    ext = [('styles.css?v=16 / site-config.js / nav-config.js / nav-dropdown.js / costing-sync.js', 'not ours', 'GitHub Pages, site root', 'Other sessions own these', 'Loaded by the page: ' + '; '.join(styles + scripts))]
    end = table(ws, heads, files + ext, [38, 12, 34, 28, 80], start=r)
    ws.freeze_panes = None

    # ================= Settings
    ws = wb.create_sheet('Settings')
    title(ws, 'Settings - v%s (%s)' % (VERSION, STAMP), 'Edit only the yellow cells. "Where" names the file and constant. Weights are auto-normalised (divided by their total), so they need not add to exactly 1.')
    S = []
    S += [
        ('PAGE', None, None, None),
        ('Analyses per browser per day', max_day, 'market-radar.js -> MAX_PER_DAY', 'Soft cap kept in the browser (localStorage key "%s"); easy to bypass. It protects shared free quotas.' % usage_key),
        ('Business types tickable in Idea check', max_ticked, 'market-radar.js -> MAX_TICKED', ''),
        ('Custom OSM types you can add', max_custom, 'market-radar.js -> MAX_CUSTOM', 'Letters, numbers, _ : - only (up to 40 characters).'),
        ('Browser-direct Overpass servers, in order', ' | '.join(direct), 'market-radar.js -> DIRECT_MIRRORS', 'Tried first, from the visitor\'s own connection; the Worker is the fallback.'),
        ('Overpass query time limit, walk areas (s)', ovp_s.group(2), 'market-radar.js -> analyze() (inline)', 'Drive areas use %s s.' % ovp_s.group(1)),
        ('Overture files folder/prefix', ovt_path, 'market-radar.js -> OVT_PATH', 'File names are <prefix><town>.json.'),
        ('Place search: minimum gap between searches (ms)', find_gap, 'market-radar.js -> findPlace() (inline)', 'Nominatim policy: at most 1 request per second per site, submit-only (no autocomplete).'),
        ('Worker address', worker_url, 'market-radar.js -> WORKER', ''),
        ('ENGINE - thresholds (all assumptions)', None, None, None),
    ]
    tnote = {
        'ANCHOR_N': 'A demand driver reaches ~63% of its pull at this many size-weighted institutions.',
        'CLUSTER_N': 'A related business type reaches ~63% of its pull at this many places.',
        'PIONEER_MIN_CONF': 'Data confidence (0-1) needed to call a missing type "First-mover" instead of "Possible gap - verify".',
        'SPARSE_N': 'With no Overture cross-check, confidence = mapped businesses / this number (max 100%).',
        'OVT_MIN_CONF': 'Overture places below this confidence are ignored.',
        'MERGE_NAMED_M': 'OSM and Overture entries with similar names within this distance are treated as the same place.',
        'MERGE_UNNAMED_M': 'Unnamed entries must be this close to be merged.',
        'MERGE_NAME_SIM': 'Share of name words that must match (0-1).',
        'COLOC_R_M': 'Radius used to test whether two business types appear together more often than chance.',
        'COLOC_MIN_N': 'Minimum count of EACH type before a co-location ratio is shown.',
        'NICHE_MIN_FOOD': 'Minimum food places before cuisine gaps are reported.',
        'NICHE_MIN_TAGGED': 'Minimum share of food places with a known cuisine before cuisine gaps are reported.',
        'THIN_SAT': 'Saturation (density / "crowded at" density) below this = under-supplied.',
        'CROWDED_SAT': 'Saturation at or above this = crowded.',
        'EMERGING_SCORE': 'National trend score (0-1) needed for the Emerging tag.',
        'CR_MIN_MATCHES': 'Minimum places found in BOTH sources before a completeness estimate is attempted.',
        'WHITE_FLOOR': 'Even a fully saturated category keeps this share of its potential (so a strong need is not zeroed out).',
    }
    for k, v in eng['T'].items():
        S.append(('Engine: ' + k, v, 'market-radar-engine.js -> T.' + k, tnote.get(k, '')))
    S.append(('ENGINE - ranking weights (assumptions)', None, None, None))
    wrow = {}
    for name, key, desc in (('Best opportunities', 'opp', 'OPP_WEIGHTS'), ('Best complements', 'comp', 'COMP_WEIGHTS'), ('Idea check', 'idea', 'IDEA_WEIGHTS')):
        first = len(S) + 5
        for k, v in eng[key].items():
            S.append(('%s: %s' % (name, k), v, 'market-radar-engine.js -> %s.%s' % (desc, k), ''))
        wrow[key] = (first, len(S) + 4)
        S.append(('%s: TOTAL (check)' % name, 'SUM', 'formula', 'Weights are divided by their total, so this need not be exactly 1.00.'))
    S.append(('ENGINE - reference points', None, None, None))
    S += [
        ('District population that maxes the Idea-check population score', eng['popNorm'], 'market-radar-engine.js -> POP_NORM', ''),
        ('Median household income that maxes the Idea-check income score (RM/month)', eng['incomeNorm'], 'market-radar-engine.js -> INCOME_NORM', ''),
    ]
    for k, v in eng['capital'].items():
        S.append(('Capital factor: ' + k, v, 'market-radar-engine.js -> CAPITAL_FACTOR.' + k, 'Multiplies the Best-opportunities score; judgement.'))
    for k, v in eng['need'].items():
        S.append(('Need tier value: ' + k, v, 'market-radar-engine.js -> NEED_TIER.' + k, 'How essential a category is to daily life; judgement.'))
    S.append(('WORKER (Cloudflare)', None, None, None))
    S += [
        ('Worker version', w_ver, 'market-radar-proxy-worker.js -> WORKER_VERSION', 'Bump on every deploy.'),
        ('Websites allowed to call the Worker', ' | '.join(origins), 'market-radar-proxy-worker.js -> ALLOWED_ORIGINS', 'CORS stops other websites, not curl. Origin enforcement/Turnstile is on the to-do list.'),
        ('Overpass servers tried by the Worker, in order', ' | '.join(mirrors), 'market-radar-proxy-worker.js -> OVERPASS_MIRRORS', 'Public servers are shared and rate-limited; Cloudflare addresses are sometimes blocked, which is why the browser tries first.'),
        ('Timeout per outside call (ms)', ext_ms, 'market-radar-proxy-worker.js -> EXTERNAL_CALL_TIMEOUT_MS', ''),
        ('Most OSM tag pairs accepted per request', max_pairs, 'market-radar-proxy-worker.js -> MAX_TAG_PAIRS', 'Everything the page sends is validated: the endpoint is public.'),
        ('Largest search box (degrees, each way)', max_span, 'market-radar-proxy-worker.js -> MAX_SPAN_DEG', ''),
        ('Overpass answer cache (seconds)', ttl_ovp, 'market-radar-proxy-worker.js (inline)', 'Only SUCCESSFUL answers are cached. Cloudflare guarantees the cache on custom domains only.'),
        ('Walk/drive-shape cache (seconds)', ttl_iso, 'market-radar-proxy-worker.js (inline)', ''),
        ('District population/income cache (seconds)', ttl_dem, 'market-radar-proxy-worker.js (inline)', 'Successes only.'),
        ('Gemini model for the plain-English read', gemini, 'market-radar-proxy-worker.js -> GEMINI_MODEL', 'Google refuses some Cloudflare regions; the Workers AI backstop covers that.'),
        ('Workers AI backstop model', wai, 'market-radar-proxy-worker.js -> WORKERS_AI_MODEL', 'Cloudflare retires models now and then; check its pricing/models page if it stops answering.'),
    ]
    heads = ['Setting', 'Current value', 'Your value', 'Where to change it', 'Notes']
    start = 4
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=start, column=j, value=h)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
    rr = start + 1
    for name, val, where, note in S:
        if val is None and where is None:
            c = ws.cell(row=rr, column=1, value=name)
            c.font = F_BOLD
            for j in range(1, 6):
                ws.cell(row=rr, column=j).fill = PatternFill('solid', fgColor='DDE8E2')
        elif val == 'SUM':
            key = [k for k in wrow if name.startswith({'opp': 'Best opportunities', 'comp': 'Best complements', 'idea': 'Idea check'}[k])][0]
            a, b = wrow[key]
            ws.cell(row=rr, column=1, value=name).font = F_BOLD
            c = ws.cell(row=rr, column=2, value='=SUM(B%d:B%d)' % (a, b))
            c.font, c.number_format = F_BOLD, '0.00'
            ws.cell(row=rr, column=4, value=where).font = F_BASE
            n = ws.cell(row=rr, column=5, value=note)
            n.font, n.alignment = F_BASE, WRAP
        else:
            ws.cell(row=rr, column=1, value=name).font = F_BASE
            c = ws.cell(row=rr, column=2, value=val)
            c.font, c.fill = F_BASE, FILL_GEN
            y = ws.cell(row=rr, column=3)
            y.fill = FILL_YOU
            ws.cell(row=rr, column=4, value=where).font = F_BASE
            n = ws.cell(row=rr, column=5, value=note)
            n.font, n.alignment = F_BASE, WRAP
        for j in range(1, 6):
            ws.cell(row=rr, column=j).border = BORDER
            ws.cell(row=rr, column=j).alignment = WRAP
        rr += 1
    for j, w in enumerate([58, 34, 18, 52, 78], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A5'

    # ================= Categories
    ws = wb.create_sheet('Categories')
    title(ws, 'Categories - v%s (%s)' % (VERSION, STAMP), '%d business types in %d groups. Crowded-at density, need tier and capital are judgements (assumptions). MSIC codes are reference only; where flagged they were cross-checked against other countries\' ISIC-derived lists, not Malaysia SSM/DOSM directly.' % (len(eng['cats']), len(eng['groups']) - 1))
    rows = []
    for c in eng['cats']:
        rows.append([c['key'], c['label'], eng['groups'][c['group']], c['needTier'], c['need'], c['crowded'], c['capital'], 'yes' if c['founder'] else 'no (left out of Best opportunities)',
                     '; '.join(c['osm']), c['ovt'], (c['msic'] or {}).get('code', ''), (c['msic'] or {}).get('name', ''), c['iowrt'] or '', '; '.join(c['geo']), None])
    table(ws, ['Key', 'Label shown on the page', 'Group', 'Need tier', 'Need value', 'Crowded at (places per km2)', 'Capital needed', 'Counts as a founder idea?', 'OSM tags queried',
               'Overture category words matched', 'MSIC 2008 code', 'MSIC name', 'DOSM trend group', 'Geoapify categories (backstop)', 'YOUR change'],
          rows, [12, 28, 20, 13, 9, 14, 11, 22, 52, 52, 10, 40, 10, 46, 26], you_cols=(15,))

    # ================= Anchors
    ws = wb.create_sheet('Anchors')
    title(ws, 'Anchors - v%s (%s)' % (VERSION, STAMP), 'Foot-traffic institutions. They are never counted as competitors; they feed the demand-driver signal. Size weight = how much one counts (assumption).')
    rows = [[a['key'], a['label'], '; '.join(a['osm']), a['ovt'], eng['anchorSize'].get(a['key']), None] for a in eng['anchors']]
    table(ws, ['Key', 'Label', 'OSM tags queried', 'Overture category words matched', 'Size weight', 'YOUR change'], rows, [12, 30, 44, 60, 12, 26], you_cols=(6,))

    # ================= Affinities
    ws = wb.create_sheet('AnchorPull')
    title(ws, 'AnchorPull - v%s (%s)' % (VERSION, STAMP), 'How strongly each anchor type draws each business type (0-1). Judgement priors, not measurements.')
    rows = []
    label = {c['key']: c['label'] for c in eng['cats']}
    alabel = {a['key']: a['label'] for a in eng['anchors']}
    for a, d in eng['anchorAffinity'].items():
        for k, w in sorted(d.items(), key=lambda x: -x[1]):
            rows.append([a, alabel.get(a, a), k, label.get(k, k), w, None])
    table(ws, ['Anchor key', 'Anchor', 'Business key', 'Business type', 'Pull (0-1)', 'YOUR change'], rows, [12, 30, 14, 30, 12, 26], you_cols=(6,))
    ws = wb.create_sheet('GoesWith')
    title(ws, 'GoesWith - v%s (%s)' % (VERSION, STAMP), 'Business types that tend to cluster (symmetric, 0-1). Judgement priors; the page also shows how often the pair is ACTUALLY seen together in the analysed area.')
    rows = [[a, label.get(a, a), b, label.get(b, b), w, None] for a, b, w in eng['pairs']]
    table(ws, ['Type A', 'Label A', 'Type B', 'Label B', 'Weight (0-1)', 'YOUR change'], rows, [14, 28, 14, 28, 12, 26], you_cols=(6,))

    # ================= Niches + hints
    ws = wb.create_sheet('Niches')
    title(ws, 'Niches and name hints - v%s (%s)' % (VERSION, STAMP), 'Cuisine sub-segments used to spot a missing niche inside a crowded food category (only reported when enough food places have a known cuisine).')
    rows = [[n['key'], n['label'], '; '.join(n['osm']), None] for n in eng['niches']]
    end = table(ws, ['Key', 'Label', 'OSM cuisine values', 'YOUR change'], rows, [14, 26, 70, 26], you_cols=(4,))
    r0 = end + 3
    ws.cell(row=r0 - 1, column=1, value='Brand/keyword list that re-classifies a place as Drink stall / bubble tea (matched anywhere in the name, lowercase)').font = F_BOLD
    rows = [[i + 1, h, None] for i, h in enumerate(eng['drinkHints'])]
    for j, h in enumerate(['#', 'Keyword', 'YOUR change / addition'], 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = F_BASE, BORDER
            if j == 3:
                c.fill = FILL_YOU
    ws.freeze_panes = 'A5'

    # ================= Catchments + Towns
    ws = wb.create_sheet('Catchments')
    title(ws, 'Catchments - v%s (%s)' % (VERSION, STAMP), 'The six walk/drive modes. The fallback circle is used only when the real travel-time shape is unavailable (no ORS key or ORS down); the page says which it used.')
    rows = [[k, l, p, int(s), int(f), None] for k, l, p, s, f in catch]
    table(ws, ['Mode key', 'Label', 'OpenRouteService profile', 'Travel time (seconds)', 'Fallback circle radius NOW (m)', 'YOUR fallback radius (m)'], rows, [12, 20, 24, 16, 22, 22], you_cols=(6,))
    ws = wb.create_sheet('Towns')
    title(ws, 'Towns - v%s (%s)' % (VERSION, STAMP), 'Adding a 5th town needs: a row in market-radar.js TOWNS, a pinned population in the Worker (DISTRICT_POPULATION_FALLBACK_2023), an entry in tools/market-radar-build-overture.py TOWNS, and a new Overture file.')
    popd = dict(pop)
    rows = []
    for k, l, la, lo, d in towns:
        bb = mrb.bbox_for(k)
        rows.append([k, l, float(la), float(lo), d, int(popd.get(k, 0)), 'W %.4f, S %.4f, E %.4f, N %.4f' % tuple(bb), None])
    table(ws, ['Key', 'Town label', 'Latitude', 'Longitude', 'DOSM district (exact spelling)', 'Pinned population (DOSM 2023 estimate)', 'Overture file covers', 'YOUR change'], rows, [12, 14, 12, 12, 24, 26, 46, 24], you_cols=(8,))
    base = 4 + len(rows) + 2
    for i, t in enumerate([
        'Pinned populations are used only when both DOSM live endpoints fail (confirmed broken on 2026-09-22). Refresh them about once a year.',
        'Overture file area: centre +/- %s degrees (about 22 km); the builder keeps places with confidence >= %s and refuses to write fewer than %d rows.' % (mrb.HALF_SPAN_DEG, mrb.MIN_CONF, mrb.MIN_ROWS),
    ]):
        ws.cell(row=base + i, column=1, value=t).font = F_NOTE

    # ================= Cloudflare
    ws = wb.create_sheet('Cloudflare')
    title(ws, 'Cloudflare - v%s (%s)' % (VERSION, STAMP), 'Secret / binding NAMES only - never values. Names must match exactly (a misspelt name is silently ignored). Names below are read from the Worker code.')
    info = {
        'GEMINI_API_KEY': ('Secret', 'Plain-English read (first choice)', 'Falls back to Workers AI if bound, else an error message', 'The same key your other tools use'),
        'ORS_API_KEY': ('Secret', 'Real walk/drive-time shape', 'The page draws a dashed circle instead', 'account.heigit.org -> "Basic Key"'),
        'GEOAPIFY_API_KEY': ('Secret', 'Backstop when the browser AND every Overpass server fail', 'Total outage shows an honest error', 'myprojects.geoapify.com (free, no card)'),
        'AI': ('Binding (Workers AI)', 'Plain-English read backstop', 'No backstop if Gemini refuses a region', 'Worker > Settings > Bindings > Add > Workers AI, variable name exactly AI'),
    }
    rows = []
    for n in env_names:
        k = info.get(n, ('?', 'Used by the Worker (not documented here yet)', '', ''))
        rows.append([n, k[0], k[1], k[2], k[3], None])
    rows.append(['GOOGLE_PLACES_API_KEY', 'Secret (no longer used)', 'Removed in v1.7.0 (it was off and paid). You may delete it from Cloudflare.', '-', '-', None])
    table(ws, ['Name (type exactly)', 'Kind', 'Needed for', 'If absent', 'Where to get it', 'Configured? (Y/N - you fill)'], rows, [26, 22, 44, 46, 50, 22], you_cols=(6,))

    wb.save(OUT)
    return OUT, [x for x in env_names]


if __name__ == '__main__':
    out, env = build()
    print('wrote', out)
    print('Worker reads these env names:', ', '.join(env))
